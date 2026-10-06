import random
import re
import os
from copy import deepcopy
import docx
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side


def setup_page_margins(doc):
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(1.2)
        section.bottom_margin = Cm(1.2)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(1.2)
        section.gutter = Cm(0.0)
        section.different_first_page_header_footer = False


def setup_document_font(doc):
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:ascii'), 'Times New Roman')
    rFonts.set(qn('w:hAnsi'), 'Times New Roman')
    rFonts.set(qn('w:cs'), 'Times New Roman')


def set_run_font_and_size(r_node, font_name="Times New Roman", size_pt=12):
    rPr = r_node.find(qn('w:rPr'))
    if rPr is None:
        rPr = OxmlElement('w:rPr')
        r_node.insert(0, rPr)
    
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    rFonts.set(qn('w:cs'), font_name)
    
    sz = rPr.find(qn('w:sz'))
    if sz is None:
        sz = OxmlElement('w:sz')
        rPr.append(sz)
    sz.set(qn('w:val'), str(int(size_pt * 2)))


def strip_all_underline_and_color(element):
    """Xóa gạch chân đáp án đúng và màu đỏ khi xuất bản đề thi cho học sinh"""
    for u in element.findall(f".//{qn('w:u')}"):
        p = u.getparent()
        if p is not None:
            p.remove(u)
    for c in element.findall(f".//{qn('w:color')}"):
        val = c.get(qn('w:val'))
        if val and val.upper() in ["FF0000", "RED", "C00000"]:
            p = c.getparent()
            if p is not None:
                p.remove(c)
    for h in element.findall(f".//{qn('w:highlight')}"):
        p = h.getparent()
        if p is not None:
            p.remove(h)


def remove_table_borders(table):
    tblPr = table._tbl.tblPr
    tblBorders = tblPr.first_child_found_in("w:tblBorders")
    if tblBorders is None:
        tblBorders = OxmlElement('w:tblBorders')
        tblPr.append(tblBorders)
    for border_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), 'none')
        tblBorders.append(border)


def apply_paragraph_format(p):
    p.style = 'Normal'
    pPr = p._p.get_or_add_pPr()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    
    for old_ind in pPr.findall(qn('w:ind')):
        pPr.remove(old_ind)
    ind = OxmlElement('w:ind')
    ind.set(qn('w:left'), '0')
    ind.set(qn('w:right'), '0')
    ind.set(qn('w:firstLine'), '0')
    pPr.append(ind)
    
    for old_sp in pPr.findall(qn('w:spacing')):
        pPr.remove(old_sp)
    spacing = OxmlElement('w:spacing')
    spacing.set(qn('w:before'), '0')
    spacing.set(qn('w:after'), '0')
    spacing.set(qn('w:line'), '340')
    spacing.set(qn('w:lineRule'), 'atLeast')
    pPr.append(spacing)
    
    if pPr.find(qn('w:contextualSpacing')) is None:
        pPr.append(OxmlElement('w:contextualSpacing'))


def set_all_left_tabs(paragraph, positions_cm):
    pPr = paragraph._p.get_or_add_pPr()
    tabs = pPr.find(qn('w:tabs'))
    if tabs is not None:
        pPr.remove(tabs)
    tabs = OxmlElement('w:tabs')
    for pos in positions_cm:
        tab = OxmlElement('w:tab')
        tab.set(qn('w:val'), 'left')
        tab.set(qn('w:leader'), 'none')
        tab.set(qn('w:pos'), str(int(round(pos * 567))))
        tabs.append(tab)
    pPr.append(tabs)


def add_tab_character(paragraph):
    r = OxmlElement('w:r')
    t = OxmlElement('w:tab')
    r.append(t)
    paragraph._p.append(r)


def append_clean_content_no_tabs(target_p, source_xml):
    for child in list(source_xml):
        tag = child.tag.split('}')[-1]
        if tag in ['pPr', 'tab']:
            continue
        new_child = deepcopy(child)
        for old_tab in new_child.findall(f".//{qn('w:tab')}"):
            p_tab = old_tab.getparent()
            if p_tab is not None:
                p_tab.remove(old_tab)
                
        if tag == 'r':
            set_run_font_and_size(new_child, "Times New Roman", 12)
        elif tag in ['oMath', 'oMathPara']:
            for r_in_math in new_child.findall(f".//{qn('w:r')}"):
                set_run_font_and_size(r_in_math, "Times New Roman", 12)
        target_p._p.append(new_child)


def append_question_body_clean(target_p, source_xml, new_q_num):
    strip_all_underline_and_color(source_xml)
    
    # 1. Tạo nhãn câu hỏi mới: "Câu x. "
    q_run = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), 'Times New Roman')
    rFonts.set(qn('w:hAnsi'), 'Times New Roman')
    rFonts.set(qn('w:cs'), 'Times New Roman')
    rPr.append(rFonts)
    sz = OxmlElement('w:sz')
    sz.set(qn('w:val'), '24')
    rPr.append(sz)
    b = OxmlElement('w:b')
    rPr.append(b)
    q_run.append(rPr)
    t_node = OxmlElement('w:t')
    t_node.text = f"Câu {new_q_num}. "
    t_node.set(qn('xml:space'), 'preserve')
    q_run.append(t_node)
    target_p._p.append(q_run)

    # 2. Xác định độ dài cụm từ "Câu [0-9]+[\.:\s]*" ở đầu câu
    full_text = "".join([t.text for t in source_xml.findall(f".//{qn('w:t')}") if t.text])
    match = re.search(r"^(?:\s*(?:Câu|Question)\s*\d+[\.:\s]*)", full_text, re.IGNORECASE)
    chars_to_remove = len(match.group(0)) if match else 0

    # 3. Duyệt qua TẤT CẢ các thành phần con trực tiếp (Text, Math, Drawing, MathType)
    for child in list(source_xml):
        tag = child.tag.split('}')[-1]
        if tag in ['pPr']:
            continue
        
        new_child = deepcopy(child)

        # Xóa nhãn "Câu x." ở các text node đầu tiên
        if chars_to_remove > 0:
            for t in new_child.findall(f".//{qn('w:t')}"):
                if not t.text:
                    continue
                if len(t.text) <= chars_to_remove:
                    chars_to_remove -= len(t.text)
                    t.text = ""
                else:
                    t.text = t.text[chars_to_remove:]
                    chars_to_remove = 0
                    break

        # Chuẩn hóa font chữ Times New Roman 12pt cho các text run
        if tag == 'r':
            set_run_font_and_size(new_child, "Times New Roman", 12)
            # Nếu run không còn chữ sau khi xóa nhãn và không chứa drawing/math thì bỏ qua
            has_text = any(t.text for t in new_child.findall(qn('w:t')))
            has_other = len(new_child.findall(f".//{qn('w:drawing')}")) > 0 or len(new_child.findall(f".//{qn('w:object')}")) > 0
            if has_text or has_other:
                target_p._p.append(new_child)
        else:
            # GIỮ NGUYÊN vẹn công thức Toán (oMath, oMathPara), MathType (object, pict) và ảnh
            target_p._p.append(new_child)

    # Dọn dẹp khoảng trắng thừa cuối đoạn
    all_runs = target_p._p.findall(qn('w:r'))
    if all_runs:
        last_run = all_runs[-1]
        for t in last_run.findall(qn('w:t')):
            if t.text:
                t.text = t.text.rstrip('\r\n ')


def bold_and_clean_any_label(p_element, new_label=None, is_letter_dot=True):
    strip_all_underline_and_color(p_element)
    runs = p_element.findall(qn('w:r'))
    if not runs:
        return

    full_text = "".join([t.text for t in p_element.findall(f".//{qn('w:t')}") if t.text])
    match = re.match(r"^(\s*)(#?[A-Da-d][\.:\)])(\s*)", full_text)
    if not match:
        return

    chars_to_remove = len(match.group(0))
    sep = "." if is_letter_dot else ")"
    old_label_str = match.groups()[1]
    target_label = f"{new_label}{sep} " if new_label is not None else f"{old_label_str.replace('#', '')} "

    for r in runs:
        for t in r.findall(qn('w:t')):
            if not t.text:
                continue
            if chars_to_remove <= 0:
                break
            if len(t.text) <= chars_to_remove:
                chars_to_remove -= len(t.text)
                t.text = ""
            else:
                t.text = t.text[chars_to_remove:].lstrip()
                chars_to_remove = 0
        if chars_to_remove <= 0:
            break

    label_r = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), 'Times New Roman')
    rFonts.set(qn('w:hAnsi'), 'Times New Roman')
    rFonts.set(qn('w:cs'), 'Times New Roman')
    rPr.append(rFonts)
    sz = OxmlElement('w:sz')
    sz.set(qn('w:val'), '24')
    rPr.append(sz)
    b = OxmlElement('w:b')
    rPr.append(b)
    label_r.append(rPr)
    label_t = OxmlElement('w:t')
    label_t.text = target_label
    label_t.set(qn('xml:space'), 'preserve')
    label_r.append(label_t)

    pPr = p_element.find(qn('w:pPr'))
    if pPr is not None:
        pPr.addnext(label_r)
    else:
        p_element.insert(0, label_r)


def get_choice_effective_length(c):
    if "text" in c and c["text"]:
        clean_t = re.sub(r"^[A-Da-d][\.:\)]\s*", "", c["text"].strip())
        return len(clean_t)
    total = 0
    for t in c["xml"].findall(f".//{qn('w:t')}"):
        if t.text:
            total += len(t.text.strip())
    return max(0, total - 3)


def insert_exam_header(doc, exam_code, exam_info=None):
    if exam_info is None:
        exam_info = {}

    so_gd = exam_info.get("so_gd") or "SỞ GD&ĐT TP CẦN THƠ"
    truong = exam_info.get("truong") or "TRƯỜNG THCS VÀ THPT QUỐC TẾ HÒA BÌNH"
    ki_thi = exam_info.get("ki_thi") or "KIỂM TRA GIỮA KỲ I"
    nam_hoc = exam_info.get("nam_hoc") or "NĂM HỌC 2026 - 2027"
    mon_hoc = exam_info.get("mon_hoc") or "MÔN: TOÁN - LỚP 12"
    thoi_gian = exam_info.get("thoi_gian") or "50 phút"

    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    remove_table_borders(table)

    table.columns[0].width = Cm(8.5)
    table.columns[1].width = Cm(9.3)

    cell_l = table.cell(0, 0)
    p_l = cell_l.paragraphs[0]
    p_l.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_l.paragraph_format.line_spacing = 1.15
    p_l.paragraph_format.space_before = Pt(0)
    p_l.paragraph_format.space_after = Pt(2)

    r_so = p_l.add_run(f"{so_gd.strip().upper()}\n")
    r_so.font.name = 'Times New Roman'
    r_so.font.size = Pt(11)
    r_so.font.color.rgb = RGBColor(220, 0, 0)

    r_tr = p_l.add_run(f"{truong.strip().upper()}")
    r_tr.bold = True
    r_tr.font.name = 'Times New Roman'
    r_tr.font.size = Pt(11.5)
    r_tr.font.color.rgb = RGBColor(220, 0, 0)

    p_box = cell_l.add_paragraph()
    p_box.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_box.paragraph_format.space_before = Pt(3)
    p_box.paragraph_format.space_after = Pt(1)
    p_box.paragraph_format.line_spacing = 1.0
    p_box.paragraph_format.left_indent = Cm(2.5)
    p_box.paragraph_format.right_indent = Cm(2.5)

    pPr = p_box._p.get_or_add_pPr()
    pBdr = parse_xml(
        f'<w:pBdr {nsdecls("w")}>\n'
        f'  <w:top w:val="single" w:sz="10" w:space="2" w:color="004080"/>\n'
        f'  <w:left w:val="single" w:sz="10" w:space="6" w:color="004080"/>\n'
        f'  <w:bottom w:val="single" w:sz="10" w:space="2" w:color="004080"/>\n'
        f'  <w:right w:val="single" w:sz="10" w:space="6" w:color="004080"/>\n'
        f'</w:pBdr>'
    )
    pPr.append(pBdr)

    r_box = p_box.add_run("ĐỀ CHÍNH THỨC")
    r_box.bold = True
    r_box.font.name = 'Times New Roman'
    r_box.font.size = Pt(10.5)
    r_box.font.color.rgb = RGBColor(220, 0, 0)

    p_sub_page = cell_l.add_paragraph()
    p_sub_page.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub_page.paragraph_format.space_before = Pt(1)
    p_sub_page.paragraph_format.space_after = Pt(0)
    p_sub_page.paragraph_format.line_spacing = 1.0
    r_sub_page = p_sub_page.add_run("(Đề thi có 04 trang)")
    r_sub_page.italic = True
    r_sub_page.font.name = 'Times New Roman'
    r_sub_page.font.size = Pt(10)

    cell_r = table.cell(0, 1)
    p_r = cell_r.paragraphs[0]
    p_r.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_r.paragraph_format.line_spacing = 1.15
    p_r.paragraph_format.space_before = Pt(0)
    p_r.paragraph_format.space_after = Pt(0)

    r_ki = p_r.add_run(f"{ki_thi.strip().upper()}\n")
    r_ki.bold = True
    r_ki.font.name = 'Times New Roman'
    r_ki.font.size = Pt(11)

    r_nam = p_r.add_run(f"{nam_hoc.strip().upper()}\n")
    r_nam.bold = True
    r_nam.font.name = 'Times New Roman'
    r_nam.font.size = Pt(11)

    mon_text = mon_hoc.strip().upper()
    if not mon_text.startswith("MÔN"):
        mon_text = f"MÔN: {mon_text}"
    r_mon = p_r.add_run(f"{mon_text}\n")
    r_mon.bold = True
    r_mon.font.name = 'Times New Roman'
    r_mon.font.size = Pt(12)

    r_time = p_r.add_run(f"Thời gian làm bài: {thoi_gian.strip()}\n(không kể thời gian phát đề)")
    r_time.italic = True
    r_time.font.name = 'Times New Roman'
    r_time.font.size = Pt(10)

    p_info = doc.add_paragraph()
    apply_paragraph_format(p_info)
    p_info.paragraph_format.space_before = Pt(8)
    p_info.paragraph_format.space_after = Pt(4)
    set_all_left_tabs(p_info, [11.0, 14.5])
    
    r_name = p_info.add_run("Họ và tên: ..............................................................")
    r_name.font.name = 'Times New Roman'
    r_name.font.size = Pt(11)
    
    add_tab_character(p_info)
    r_sbd = p_info.add_run("Số báo danh: .......")
    r_sbd.font.name = 'Times New Roman'
    r_sbd.font.size = Pt(11)
    
    add_tab_character(p_info)
    r_code = p_info.add_run(f"Mã đề {exam_code}")
    r_code.bold = True
    r_code.font.name = 'Times New Roman'
    r_code.font.size = Pt(12)

    pPr_info = p_info._p.get_or_add_pPr()
    pBdr_info = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="12" w:space="3" w:color="000000"/></w:pBdr>')
    pPr_info.append(pBdr_info)


def add_field_code(paragraph, field_text):
    run = paragraph.add_run()
    r = run._r
    fldSimple = OxmlElement('w:fldSimple')
    fldSimple.set(qn('w:instr'), field_text)
    r.append(fldSimple)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(10)


def setup_exam_footer(doc, exam_code):
    for section in doc.sections:
        footer = section.footer
        p_ft = footer.paragraphs[0]
        p_ft.alignment = WD_ALIGN_PARAGRAPH.LEFT
        
        pPr = p_ft._p.get_or_add_pPr()
        for old_bdr in pPr.findall(qn('w:pBdr')):
            pPr.remove(old_bdr)
        pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:top w:val="single" w:sz="4" w:space="3" w:color="888888"/></w:pBdr>')
        pPr.append(pBdr)

        set_all_left_tabs(p_ft, [17.5])
        
        r_code = p_ft.add_run(f"Mã đề {exam_code}")
        r_code.font.name = 'Times New Roman'
        r_code.font.size = Pt(10)
        r_code.font.color.rgb = RGBColor(100, 100, 100)

        add_tab_character(p_ft)

        r_pg = p_ft.add_run("Trang ")
        r_pg.font.name = 'Times New Roman'
        r_pg.font.size = Pt(10)
        r_pg.font.color.rgb = RGBColor(100, 100, 100)

        add_field_code(p_ft, "PAGE")
        
        r_slash = p_ft.add_run("/")
        r_slash.font.name = 'Times New Roman'
        r_slash.font.size = Pt(10)
        r_slash.font.color.rgb = RGBColor(100, 100, 100)

        add_field_code(p_ft, "NUMPAGES")


def insert_choices_smart(doc, choices, is_letter_dot=True):
    if not choices:
        return

    lengths = [get_choice_effective_length(c) for c in choices]
    max_len = max(lengths) if lengths else 0

    if len(choices) == 4 and max_len <= 12:
        row_p = doc.add_paragraph()
        apply_paragraph_format(row_p)
        set_all_left_tabs(row_p, [0.5, 5.13, 9.75, 14.37])
        for c in choices:
            add_tab_character(row_p)
            append_clean_content_no_tabs(row_p, c["xml"])

    elif len(choices) == 4 and max_len <= 38:
        r1 = doc.add_paragraph()
        apply_paragraph_format(r1)
        set_all_left_tabs(r1, [0.5, 9.5])
        add_tab_character(r1)
        append_clean_content_no_tabs(r1, choices[0]["xml"])
        add_tab_character(r1)
        append_clean_content_no_tabs(r1, choices[1]["xml"])
        
        r2 = doc.add_paragraph()
        apply_paragraph_format(r2)
        set_all_left_tabs(r2, [0.5, 9.5])
        add_tab_character(r2)
        append_clean_content_no_tabs(r2, choices[2]["xml"])
        add_tab_character(r2)
        append_clean_content_no_tabs(r2, choices[3]["xml"])

    else:
        for c in choices:
            p = doc.add_paragraph()
            apply_paragraph_format(p)
            set_all_left_tabs(p, [0.5])
            add_tab_character(p)
            append_clean_content_no_tabs(p, c["xml"])


def is_choice_correct(choice_obj):
    if choice_obj.get("is_correct") is True:
        return True
    
    xml_node = choice_obj.get("xml")
    if xml_node is not None:
        for u in xml_node.findall(f".//{qn('w:u')}"):
            val = u.get(qn('w:val'))
            if val is None or val not in ['none', '0']:
                return True
        for col in xml_node.findall(f".//{qn('w:color')}"):
            val = col.get(qn('w:val'))
            if val and val.upper() in ["FF0000", "RED", "C00000"]:
                return True
        for h in xml_node.findall(f".//{qn('w:highlight')}"):
            val = h.get(qn('w:val'))
            if val and val not in ['none']:
                return True

    text_val = choice_obj.get("text", "")
    if text_val and (text_val.strip().startswith("#") or text_val.strip().startswith("*")):
        return True

    return False


def is_element_containing_header_rubbish(elem):
    texts = [t.text for t in elem.findall(f".//{qn('w:t')}") if t.text]
    full_str = " ".join(texts).upper()
    keywords = ["ĐỀ CHÍNH THỨC", "DE CHINH THUC", "ĐỀ THI CÓ", "DE THI CO", "SỞ GD&ĐT", "SO GD&DT", "SỐ BÁO DANH", "HỌ VÀ TÊN"]
    return any(kw in full_str for kw in keywords)


def shuffle_and_export(groups, original_doc_path, exam_code, output_path, exam_info=None):
    rng = random.Random(str(exam_code))
    
    # Mở trực tiếp từ file gốc để bảo toàn toàn bộ image part, font, styles và relationship
    doc = docx.Document(original_doc_path)
    
    # Xóa sạch toàn bộ body cũ, giữ lại các cấu hình ngầm của file Word
    body_elem = doc._body._element
    for child in list(body_elem):
        if child.tag.endswith('sectPr'):
            continue
        body_elem.remove(child)

    setup_page_margins(doc)
    setup_document_font(doc)

    insert_exam_header(doc, exam_code, exam_info)
    setup_exam_footer(doc, exam_code)

    structured_answers = {
        "part_1": {},
        "part_2": {},
        "part_3": {}
    }
    flat_answer_keys = {}
    total_q_counter = 1

    for grp in groups:
        grp_title = grp.get("title", "")
        part_type = grp.get("part_type", "part_1")

        is_q0 = (not grp.get("shuffle_questions", True)) and (not grp.get("shuffle_choices", True))

        if grp_title:
            p_grp = doc.add_paragraph()
            apply_paragraph_format(p_grp)
            p_grp.paragraph_format.space_before = Pt(6)
            p_grp.paragraph_format.space_after = Pt(2)
            r_t = p_grp.add_run(grp_title)
            r_t.bold = True
            r_t.font.name = 'Times New Roman'
            r_t.font.size = Pt(12)

        questions = [deepcopy(q) for q in grp["questions"]]
        if grp.get("shuffle_questions", True):
            rng.shuffle(questions)

        group_local_counter = 1

        for q in questions:
            display_num = group_local_counter

            # 1. Thân câu hỏi
            p_body = doc.add_paragraph()
            apply_paragraph_format(p_body)
            append_question_body_clean(p_body, q["body_elements"][0], display_num)

            # 2. Đoạn phụ, bảng biểu, hình ảnh
            for extra in q["body_elements"][1:]:
                if is_element_containing_header_rubbish(extra):
                    continue
                extra_copy = deepcopy(extra)
                bold_and_clean_any_label(extra_copy)
                
                tag_name = extra_copy.tag.split('}')[-1]
                if tag_name == 'tbl':
                    # Gắn thẻ Table đúng chuẩn XML vào body
                    sectPr = doc._body._element.xpath('w:sectPr')
                    if sectPr:
                        sectPr[0].addprevious(extra_copy)
                    else:
                        doc._body._element.append(extra_copy)
                else:
                    p_extra = doc.add_paragraph()
                    p_extra._p.getparent().replace(p_extra._p, extra_copy)
                    apply_paragraph_format(p_extra)
                    for r in extra_copy.findall(f".//{qn('w:r')}"):
                        set_run_font_and_size(r, "Times New Roman", 12)

            # 3. Các đáp án lựa chọn
            choices = q.get("choices", [])
            if choices:
                if part_type == "part_3":
                    ans_val = choices[0].get("text", "").strip()
                    ans_val = re.sub(r"<[^>]+>", "", ans_val).strip()
                    if not is_q0:
                        structured_answers["part_3"][group_local_counter] = ans_val
                        flat_answer_keys[total_q_counter] = ans_val

                elif part_type == "part_2":
                    labels = ["a", "b", "c", "d"]
                    for idx, c in enumerate(choices):
                        lbl = labels[idx] if idx < len(labels) else c.get("label", "")
                        is_corr = is_choice_correct(c)
                        ans_val = "Đ" if is_corr else "S"
                        
                        if not is_q0:
                            key_name = f"{group_local_counter}{lbl}"
                            structured_answers["part_2"][key_name] = ans_val
                            flat_answer_keys[total_q_counter] = ans_val
                        
                        bold_and_clean_any_label(c["xml"], new_label=lbl, is_letter_dot=False)

                    insert_choices_smart(doc, choices, is_letter_dot=False)

                else:
                    labels = ["A", "B", "C", "D"]
                    for c in choices:
                        c["_detected_correct"] = is_choice_correct(c)

                    if grp.get("shuffle_choices", True):
                        unfixed = [c for c in choices if not c.get("is_fixed", False)]
                        rng.shuffle(unfixed)
                        new_choices = []
                        u_idx = 0
                        for c in choices:
                            if c.get("is_fixed", False):
                                new_choices.append(c)
                            else:
                                new_choices.append(unfixed[u_idx])
                                u_idx += 1
                        choices = new_choices

                    for idx, c in enumerate(choices):
                        lbl = labels[idx] if idx < len(labels) else c.get("label", "")
                        if c.get("_detected_correct") and not is_q0:
                            structured_answers["part_1"][group_local_counter] = lbl
                            flat_answer_keys[total_q_counter] = lbl
                            
                        bold_and_clean_any_label(c["xml"], new_label=lbl, is_letter_dot=True)

                    insert_choices_smart(doc, choices, is_letter_dot=True)

            group_local_counter += 1
            if not is_q0:
                total_q_counter += 1

    p_end = doc.add_paragraph()
    apply_paragraph_format(p_end)
    p_end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_end.paragraph_format.space_before = Pt(8)
    p_end.add_run("------ HẾT ------").bold = True

    doc.save(output_path)
    return flat_answer_keys, structured_answers


def export_answers_excel(all_parts_data, output_excel_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Dap_An"
    ws.views.sheetView[0].showGridLines = True

    thin_border = Border(
        left=Side(style='thin', color='A0A0A0'),
        right=Side(style='thin', color='A0A0A0'),
        top=Side(style='thin', color='A0A0A0'),
        bottom=Side(style='thin', color='A0A0A0')
    )
    title_font = Font(name="Times New Roman", size=12, bold=True)
    header_font = Font(name="Times New Roman", size=11, bold=True)
    cell_font = Font(name="Times New Roman", size=11)
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")

    current_row = 1

    parts_config = [
        ("part_1", "PHẦN I. Câu trắc nghiệm nhiều phương án lựa chọn"),
        ("part_2", "PHẦN II. Câu trắc nghiệm đúng sai."),
        ("part_3", "PHẦN III. Câu trắc nghiệm yêu cầu trả lời ngắn.")
    ]

    for part_key, part_title in parts_config:
        data = all_parts_data.get(part_key)
        if not data or not data.get("headers"):
            continue

        ws.cell(row=current_row, column=1, value=part_title).font = title_font
        ws.cell(row=current_row, column=1).alignment = left_align
        current_row += 1

        ws.cell(row=current_row, column=1, value=r"Đề\câu").font = header_font
        ws.cell(row=current_row, column=1).alignment = center_align
        ws.cell(row=current_row, column=1).border = thin_border

        for col_idx, h in enumerate(data["headers"], start=2):
            cell = ws.cell(row=current_row, column=col_idx, value=str(h))
            cell.font = header_font
            cell.alignment = center_align
            cell.border = thin_border
        current_row += 1

        for r in data.get("rows", []):
            code_cell = ws.cell(row=current_row, column=1, value=str(r["code"]))
            code_cell.font = header_font
            code_cell.alignment = center_align
            code_cell.border = thin_border

            for col_idx, ans in enumerate(r.get("answers", []), start=2):
                c = ws.cell(row=current_row, column=col_idx, value=str(ans) if ans is not None else "")
                c.font = cell_font
                c.alignment = center_align
                c.border = thin_border
            current_row += 1

    wb.save(output_excel_path)